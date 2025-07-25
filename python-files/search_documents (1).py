import os
import fitz  # PyMuPDF
import docx
import openpyxl

# Define the folder to search
folder_path = r"C:\Users\sodhisb\Downloads\ETQ"

# Define the keywords to search for
keywords = ["PeopleSoft", "People Soft", "Part Number", "Part No."]

# Function to search keywords in PDF files
def search_pdf(file_path):
    try:
        with fitz.open(file_path) as doc:
            for page in doc:
                text = page.get_text()
                if any(keyword in text for keyword in keywords):
                    return True
    except Exception as e:
        print(f"Error reading PDF {file_path}: {e}")
    return False

# Function to search keywords in DOCX files
def search_docx(file_path):
    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            if any(keyword in para.text for keyword in keywords):
                return True
    except Exception as e:
        print(f"Error reading DOCX {file_path}: {e}")
    return False

# Function to search keywords in XLSX files
def search_xlsx(file_path):
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell and any(keyword in str(cell) for keyword in keywords):
                        return True
    except Exception as e:
        print(f"Error reading XLSX {file_path}: {e}")
    return False

# Main function to scan the directory
def search_documents():
    matching_files = []
    for root, dirs, files in os.walk(folder_path):
        for file in files:
            file_path = os.path.join(root, file)
            if file.lower().endswith(".pdf") and search_pdf(file_path):
                matching_files.append(file_path)
            elif file.lower().endswith(".docx") and search_docx(file_path):
                matching_files.append(file_path)
            elif file.lower().endswith(".xlsx") and search_xlsx(file_path):
                matching_files.append(file_path)

    print("\nFiles containing the specified keywords:")
    for f in matching_files:
        print(f)

if __name__ == "__main__":
    search_documents()
