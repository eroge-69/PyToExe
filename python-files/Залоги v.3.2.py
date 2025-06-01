import sqlite3
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import time
from datetime import datetime, timedelta
import requests
import json
from threading import Thread
from cryptography.fernet import Fernet
from tkinter import Canvas
import os
import sys
import subprocess
import pythoncom
from pytz import timezone
import logging

try:
    import ttkbootstrap as ttk
    from ttkbootstrap.constants import *
    from ttkbootstrap.tooltip import ToolTip
except ImportError:
    ttk = tk  # Fallback на стандартный tkinter, если ttkbootstrap недоступен
    ToolTip = lambda widget, text, bootstyle=None: None
    PRIMARY, DANGER, SUCCESS, INFO, SECONDARY = "primary", "danger", "success", "info", "secondary"

# Настройка логирования
logging.basicConfig(
    level=logging.DEBUG,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('pawnshop.log', encoding='utf-8'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

class PawnShopApp:
    def __init__(self, root):
        logger.info("Запуск приложения PawnShopApp")
        self.root = root
        self.root.title("Учет залогов (Синхронизированная версия)")
        self.root.geometry("1200x800")
        self.root.minsize(800, 600)

        # Инициализация настроек (перемещаем сюда)
        self.settings = {
            'cell_from': 1,
            'cell_to': 200,
            'default_types': ['Паспорт', 'Водительское удостоверение', 'Карта гостя', 'Свидетельство о регистрации транспортного средства', 'Денежные средства', 'другое'],
            'server_url': 'http://localhost:5000',
            'auto_sync': True,
            'contract_template': os.path.join(os.getcwd(), "Договоры", "шаблон_договора.docx"),
            'timezone': 'Europe/Moscow',
            'theme': 'zerom',
            'sync_interval': 5,
            'sync_on_change': False,
            'database_mode': 'local',
            'local_db_path': os.path.join(os.getcwd(), 'pawnshop.db')
        }
        logger.debug("Настройки по умолчанию загружены")

        # Инициализация шифрования
        try:
            self.init_encryption()
        except Exception as e:
            logger.error(f"Ошибка при инициализации шифрования: {e}")
            messagebox.showerror("Ошибка", "Не удалось инициализировать шифрование")
            self.root.destroy()
            return

        # Подключение к базе данных
        try:
            self.conn = self.create_connection()
            if not self.conn:
                logger.error("Не удалось подключиться к базе данных")
                messagebox.showerror("Ошибка", "Не удалось подключиться к базе данных")
                self.root.destroy()
                return
            logger.info("Подключение к базе данных успешно")
        except Exception as e:
            logger.error(f"Ошибка при подключении к базе данных: {e}")
            messagebox.showerror("Ошибка", f"Ошибка базы данных: {e}")
            self.root.destroy()
            return

        # Попытка загрузки темы
        try:
            saved_theme = self.get_setting('theme', 'zerom')
            if 'ttkbootstrap' in sys.modules:
                self.style = ttk.Style(saved_theme)
                self.root.configure(bg=self.style.colors.bg)
                logger.info(f"Тема успешно установлена: {saved_theme}")
            else:
                logger.warning("ttkbootstrap не установлен, используется стандартный tkinter")
                self.style = ttk.Style()
        except Exception as e:
            logger.error(f"Ошибка при установке темы: {e}")
            self.style = ttk.Style()

        # Создание интерфейса
        try:
            self.create_widgets()
            self.update_combobox()
            self.setup_hotkeys()
            logger.info("Интерфейс успешно создан")
        except Exception as e:
            logger.error(f"Ошибка при создании интерфейса: {e}")
            messagebox.showerror("Ошибка", f"Ошибка создания интерфейса: {e}")
            self.root.destroy()
            return

        # Первая синхронизация при запуске
        try:
            if self.get_setting('auto_sync', 'True') == 'True':
                self.auto_sync_enabled = True
                self.auto_sync_enabled = self.get_setting('auto_sync', 'True') == 'True'
                self.sync_data()
                logger.info("Автосинхронизация включена и запущена")
            else:
                self.auto_sync_enabled = False
                logger.info("Автосинхронизация отключена")
        except Exception as e:
            logger.error(f"Ошибка при первой синхронизации: {e}")

        # Автосинхронизация каждые 5 минут
        self.root.after(300000, self.auto_sync)
        logger.debug("Запланирована автосинхронизация каждые 5 минут")

    def create_scrollable_frame(self, parent):
        logger.debug("Создание прокручиваемого фрейма")
        # Создаем контейнер для Canvas и полос прокрутки
        container = ttk.Frame(parent)
        container.grid(row=0, column=0, sticky="nsew")
        parent.grid_rowconfigure(0, weight=1)  # Важно: растягиваем строку родителя
        parent.grid_columnconfigure(0, weight=1)  # Важно: растягиваем столбец родителя

        container.grid_rowconfigure(0, weight=1)
        container.grid_columnconfigure(0, weight=1)

        canvas = Canvas(container, highlightthickness=0)
        scrollbar_y = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
        scrollbar_x = ttk.Scrollbar(container, orient="horizontal", command=canvas.xview)
        frame = ttk.Frame(canvas)

        def configure_scrollregion(event):
            logger.debug("Обновление scrollregion")
            canvas.configure(scrollregion=canvas.bbox("all"))

        frame.bind("<Configure>", configure_scrollregion)
        canvas.create_window((0, 0), window=frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar_y.set, xscrollcommand=scrollbar_x.set)

        # Размещение с помощью grid для адаптивности
        canvas.grid(row=0, column=0, sticky="nsew")
        scrollbar_y.grid(row=0, column=1, sticky="ns")
        scrollbar_x.grid(row=1, column=0, sticky="ew")

        return frame

    def get_current_time(self):
        try:
            tz = timezone(self.get_setting('timezone', 'Europe/Moscow'))
            current_time = datetime.now(tz)
            logger.debug(f"Текущее время: {current_time}")
            return current_time
        except Exception as e:
            logger.error(f"Ошибка получения времени: {e}")
            return datetime.now()

    def init_encryption(self):
        logger.debug("Инициализация шифрования")
        key_file = 'secret.key'
        try:
            if not os.path.exists(key_file):
                key = Fernet.generate_key()
                with open(key_file, 'wb') as f:
                    f.write(key)
                logger.info("Создан новый ключ шифрования")
            with open(key_file, 'rb') as f:
                self.cipher = Fernet(f.read())
            logger.debug("Шифрование успешно инициализировано")
        except Exception as e:
            logger.error(f"Ошибка при создании/чтении ключа шифрования: {e}")
            raise

    def update_db_settings(self):
        if self.db_mode_var.get() == 'local':
            self.local_db_frame.grid()
        else:
            self.local_db_frame.grid_remove()
        self.set_setting('database_mode', self.db_mode_var.get())

    def browse_local_db(self):
        filename = filedialog.asksaveasfilename(
            title="Укажите файл базы данных",
            filetypes=[("SQLite Database", "*.db"), ("All files", "*.*")],
            defaultextension=".db"
        )
        if filename:
            self.local_db_entry.delete(0, tk.END)
            self.local_db_entry.insert(0, filename)
            self.set_setting('local_db_path', filename)

    def delete_selected_history(self):
        selected = self.tree_history.selection()
        if not selected:
            messagebox.showwarning("Ошибка", "Выберите записи для удаления!")
            return
        
        if not messagebox.askyesno(
            "Подтверждение", 
            f"Вы действительно хотите удалить записи: {len(selected)}шт?"
        ):
            return

        try:
            cursor = self.conn.cursor()
            for item in selected:
                record_id = self.tree_history.item(item)['values'][0]
                cursor.execute('DELETE FROM history WHERE id = ?', (record_id,))
            self.conn.commit()
            self.sync_on_change()
            self.show_history()
            messagebox.showinfo("Успех", "Записи успешно удалены!")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Ошибка при удалении: {e}")

    def delete_history_by_period(self):
        period_dialog = tk.Toplevel(self.root)
        period_dialog.title("Удаление за период")
        period_dialog.geometry("400x300")
    
        # По умолчанию - вчерашний день
        yesterday = datetime.now() - timedelta(days=1)
    
        ttk.Label(period_dialog, text="Начальная дата:").pack(pady=5)
        start_frame = ttk.Frame(period_dialog)
        start_frame.pack(pady=5)
    
        self.start_day = ttk.Spinbox(start_frame, from_=1, to=31, width=5)
        self.start_day.pack(side="left", padx=5)
        self.start_day.set(yesterday.day)
    
        self.start_month = ttk.Spinbox(start_frame, from_=1, to=12, width=5)
        self.start_month.pack(side="left", padx=5)
        self.start_month.set(yesterday.month)
    
        self.start_year = ttk.Spinbox(start_frame, from_=2000, to=2100, width=7)
        self.start_year.pack(side="left", padx=5)
        self.start_year.set(yesterday.year)
    
        ttk.Label(period_dialog, text="Конечная дата:").pack(pady=5)
        end_frame = ttk.Frame(period_dialog)
        end_frame.pack(pady=5)
    
        self.end_day = ttk.Spinbox(end_frame, from_=1, to=31, width=5)
        self.end_day.pack(side="left", padx=5)
        self.end_day.set(yesterday.day)
    
        self.end_month = ttk.Spinbox(end_frame, from_=1, to=12, width=5)
        self.end_month.pack(side="left", padx=5)
        self.end_month.set(yesterday.month)
    
        self.end_year = ttk.Spinbox(end_frame, from_=2000, to=2100, width=7)
        self.end_year.pack(side="left", padx=5)
        self.end_year.set(yesterday.year)
    
        def delete_period():
            try:
                start_date = datetime(
                    int(self.start_year.get()),
                    int(self.start_month.get()),
                    int(self.start_day.get())
                )
                end_date = datetime(
                    int(self.end_year.get()),
                    int(self.end_month.get()),
                    int(self.end_day.get())
                ) + timedelta(days=1)  # Чтобы включить весь конечный день
            
                cursor = self.conn.cursor()
                cursor.execute(
                    "DELETE FROM history WHERE released_at BETWEEN ? AND ?",
                    (start_date.strftime("%Y-%m-%d %H:%M:%S"), 
                     end_date.strftime("%Y-%m-%d %H:%M:%S"))
                )
                deleted_count = cursor.rowcount
                self.conn.commit()
                self.sync_on_change()
                period_dialog.destroy()
                self.show_history()
                messagebox.showinfo("Успех", f"Удалено записей: {deleted_count}!")
            except Exception as e:
                messagebox.showerror("Ошибка", f"Ошибка при удалении: {e}")
    
        btn_frame = ttk.Frame(period_dialog)
        btn_frame.pack(pady=15)
        ttk.Button(btn_frame, text="Удалить", command=delete_period).pack(side="left", padx=10)
        ttk.Button(btn_frame, text="Отмена", command=period_dialog.destroy).pack(side="left", padx=10)
        
    def create_connection(self):
        mode = self.get_setting('database_mode', 'server')
        if mode == 'local':
            db_path = self.get_setting('local_db_path', self.settings['local_db_path'])
            try:
                conn = sqlite3.connect(db_path)
                logger.info(f"Подключение к локальной базе: {db_path}")
            except Exception as e:
                logger.error(f"Ошибка подключения к локальной базе: {e}")
                return None
        else:  # server mode
            try:
                conn = sqlite3.connect('pawnshop.db')
                logger.info("Подключение к серверной базе")
            except Exception as e:
                logger.error(f"Ошибка подключения к серверной базе: {e}")
                return None
        try:
            cursor = conn.cursor()
            cursor.execute('''
            CREATE TABLE IF NOT EXISTS clients (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                last_name TEXT NOT NULL,
                first_name TEXT NOT NULL,
                phone TEXT NOT NULL,
                item_type TEXT NOT NULL,
                item_info TEXT,
                cell_number INTEGER UNIQUE,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )''')
            cursor.execute('''
            CREATE TABLE IF NOT EXISTS history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                last_name TEXT NOT NULL,
                first_name TEXT NOT NULL,
                phone TEXT NOT NULL,
                item_type TEXT NOT NULL,
                item_info TEXT,
                cell_number INTEGER,
                created_at TEXT NOT NULL,
                released_at TEXT NOT NULL,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )''')
            cursor.execute('''
            CREATE TABLE IF NOT EXISTS cells (
                cell_number INTEGER PRIMARY KEY,
                is_occupied BOOLEAN DEFAULT FALSE
            )''')
            cursor.execute('''
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT
            )''')
            cursor.execute('SELECT value FROM settings WHERE key = "cell_from"')
            result_from = cursor.fetchone()
            cell_from = int(result_from[0]) if result_from else self.settings['cell_from']
            cursor.execute('SELECT value FROM settings WHERE key = "cell_to"')
            result_to = cursor.fetchone()
            cell_to = int(result_to[0]) if result_to else self.settings['cell_to']
            for cell in range(cell_from, cell_to + 1):
                cursor.execute('INSERT OR IGNORE INTO cells (cell_number) VALUES (?)', (cell,))
            for key, value in self.settings.items():
                cursor.execute('INSERT OR IGNORE INTO settings (key, value) VALUES (?, ?)', (key, str(value)))
            conn.commit()
            logger.info("База данных успешно инициализирована")
            return conn
        except sqlite3.Error as e:
            logger.error(f"Ошибка при создании базы данных: {e}")
            return None

    def create_widgets(self):
        logger.debug("Создание виджетов интерфейса")
        main_container = ttk.Frame(self.root, padding="10")
        main_container.grid(row=0, column=0, sticky="nsew")

        # Настройка веса для растяжения
        self.root.grid_rowconfigure(0, weight=1)
        self.root.grid_columnconfigure(0, weight=1)
        main_container.grid_rowconfigure(1, weight=1)  # Вторая строка (для tab_control) растягивается
        main_container.grid_columnconfigure(0, weight=1)

        header = ttk.Label(
            main_container,
            text="Учет залогов",
            font=("Helvetica", 24, "bold"),
            style="primary.TLabel" if 'ttkbootstrap' in sys.modules else None
        )
        header.grid(row=0, column=0, pady=(10, 20), sticky="ew")

        self.tab_control = ttk.Notebook(main_container, style="primary.TNotebook" if 'ttkbootstrap' in sys.modules else None)
        self.tab_add = ttk.Frame(self.tab_control)
        self.tab_search = ttk.Frame(self.tab_control)
        self.tab_history = ttk.Frame(self.tab_control)
        self.tab_settings = ttk.Frame(self.tab_control)

        # Настройка растяжения для вкладок
        self.tab_control.grid(row=1, column=0, sticky="nsew")

        # Прокручиваемые фреймы
        self.scroll_frame_add = self.create_scrollable_frame(self.tab_add)
        self.scroll_frame_search = self.create_scrollable_frame(self.tab_search)
        self.scroll_frame_history = self.create_scrollable_frame(self.tab_history)
        self.scroll_frame_settings = self.create_scrollable_frame(self.tab_settings)

        # Настройка вкладок
        self.setup_add_tab()
        self.setup_search_tab()
        self.setup_history_tab()
        self.setup_settings_tab()

        self.tab_control.add(self.tab_add, text=" Добавить клиента ")
        self.tab_control.add(self.tab_search, text=" Поиск/Удаление ")
        self.tab_control.add(self.tab_history, text=" История ")
        self.tab_control.add(self.tab_settings, text=" Настройки ")

        self.tab_control.bind("<<NotebookTabChanged>>", self.on_tab_changed)
        logger.debug("Виджеты интерфейса успешно созданы")

    def setup_add_tab(self):
        logger.debug("Настройка вкладки 'Добавить клиента'")
        frame = self.scroll_frame_add
        frame.configure(padding="20")
        frame.grid_rowconfigure(1, weight=1)
        frame.grid_rowconfigure(2, weight=0)
        frame.grid_rowconfigure(3, weight=0)
        frame.grid_columnconfigure(0, weight=1)

        ttk.Label(
            frame,
            text="Добавление клиента",
            font=("Helvetica", 16, "bold"),
            style="primary.TLabel" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=0, column=0, pady=(0, 20), sticky="n")

        fields = [
            ("Фамилия*:", "last_name"),
            ("Имя*:", "first_name"),
            ("Телефон*:", "phone"),
            ("Тип залога*:", "item_type"),
            ("Описание/Сумма:", "item_info")
        ]

        self.entries = {}
        input_frame = ttk.Frame(frame)
        input_frame.grid(row=1, column=0, sticky="nsew", pady=10)
        input_frame.grid_columnconfigure(1, weight=1)
        for i in range(len(fields)):
            input_frame.grid_rowconfigure(i, weight=1)

        for i, (label, name) in enumerate(fields):
            ttk.Label(
                input_frame,
                text=label,
                font=("Helvetica", 12)
            ).grid(row=i, column=0, padx=10, pady=5, sticky="e")

            if name == "item_type":
                entry = ttk.Combobox(
                    input_frame,
                    font=("Helvetica", 12),
                    style="primary.TCombobox" if 'ttkbootstrap' in sys.modules else None
                )
            else:
                entry = ttk.Entry(
                    input_frame,
                    font=("Helvetica", 12),
                    style="primary.TEntry" if 'ttkbootstrap' in sys.modules else None
                )
            entry.grid(row=i, column=1, padx=10, pady=5, sticky="ew")
            self.entries[name] = entry
            ToolTip(entry, text=f"Введите {label.strip(':*')}", bootstyle="inverse-primary")

        ttk.Button(
            frame,
            text="Добавить клиента",
            command=self.add_client,
            style="primary.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=10
        ).grid(row=2, column=0, pady=20, sticky="ew", padx=100)

        ttk.Label(
            frame,
            text="* - обязательные поля",
            font=("Helvetica", 10, "italic"),
            style="secondary.TLabel" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=3, column=0, sticky="s")

        logger.debug("Вкладка 'Добавить клиента' настроена")

    def setup_search_tab(self):
        logger.debug("Настройка вкладки 'Поиск/Удаление'")
        frame = self.scroll_frame_search
        frame.configure(padding="20")
        frame.grid_rowconfigure(2, weight=1)  # Таблица растягивается
        frame.grid_columnconfigure(0, weight=1)

        ttk.Label(
            frame,
            text="Поиск и управление клиентами",
            font=("Helvetica", 16, "bold"),
            style="primary.TLabel" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=0, column=0, pady=(0, 20), sticky="n")

        search_frame = ttk.Frame(frame)
        search_frame.grid(row=1, column=0, sticky="ew", pady=10)
        search_frame.grid_columnconfigure(1, weight=1)

        ttk.Label(
            search_frame,
            text="Поиск:",
            font=("Helvetica", 12)
        ).grid(row=0, column=0, padx=10, sticky="w")

        self.search_entry = ttk.Entry(
            search_frame,
            font=("Helvetica", 12),
            style="primary.TEntry" if 'ttkbootstrap' in sys.modules else None
        )
        self.search_entry.grid(row=0, column=1, padx=10, sticky="ew")
        ToolTip(self.search_entry, text="Введите фамилию, имя, телефон или номер ячейки", bootstyle="inverse-primary")

        ttk.Button(
            search_frame,
            text="Найти",
            command=self.find_client,
            style="primary.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=8
        ).grid(row=0, column=2, padx=5)

        ttk.Button(
            search_frame,
            text="Показать все",
            command=self.show_all_active,
            style="primary.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=8
        ).grid(row=0, column=3, padx=5)

        columns = [
            ("id", "ID", 50),
            ("last_name", "Фамилия", 120),
            ("first_name", "Имя", 120),
            ("phone", "Телефон", 120),
            ("item_type", "Тип залога", 120),
            ("item_info", "Описание", 150),
            ("cell_number", "Ячейка", 80),
            ("created_at", "Дата приёма", 150)
        ]

        tree_frame = ttk.Frame(frame)
        tree_frame.grid(row=2, column=0, sticky="nsew")
        tree_frame.grid_rowconfigure(0, weight=1)
        tree_frame.grid_columnconfigure(0, weight=1)

        self.tree_search = ttk.Treeview(
            tree_frame,
            columns=[col[0] for col in columns],
            show="headings",
            style="primary.Treeview" if 'ttkbootstrap' in sys.modules else None
        )

        for col in columns:
            self.tree_search.heading(col[0], text=col[1], 
                                    command=lambda c=col[0]: self.sort_treeview(self.tree_search, c))
            self.tree_search.column(col[0], width=col[2])

        def resize_columns(event):
            total_width = self.tree_search.winfo_width()
            num_columns = len(columns)
            for col, _, _ in columns:
                self.tree_search.column(col, width=max(50, int(total_width / num_columns)))

        self.tree_search.bind("<Configure>", resize_columns)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree_search.yview)
        self.tree_search.configure(yscrollcommand=scrollbar.set)
        self.tree_search.grid(row=0, column=0, sticky="nsew", pady=10)
        scrollbar.grid(row=0, column=1, sticky="ns")

        btn_frame = ttk.Frame(frame)
        btn_frame.grid(row=3, column=0, sticky="ew", pady=10)
        btn_frame.grid_columnconfigure((0, 1, 2, 3), weight=1)

        buttons = [
            ("Вернуть залог", self.delete_client, "danger-outline"),
            ("Редактировать", self.edit_client, "info-outline"),
            ("Дублировать", self.duplicate_client, "secondary-outline"),
            ("Печать договора", self.print_contract, "success-outline")
        ]

        for i, (text, command, style) in enumerate(buttons):
            style = f"{style}.TButton" if 'ttkbootstrap' in sys.modules else None
            btn = ttk.Button(
                btn_frame,
                text=text,
                command=command,
                style=style,
                padding=8
            )
            btn.grid(row=0, column=i, padx=5, sticky="ew")
            ToolTip(btn, text=text, bootstyle="inverse-primary")
        logger.debug("Вкладка 'Поиск/Удаление' настроена")
        
    def setup_history_tab(self):
        logger.debug("Настройка вкладки 'История'")
        frame = self.scroll_frame_history
        frame.configure(padding="20")
        frame.grid_rowconfigure(2, weight=1)
        frame.grid_columnconfigure(0, weight=1)

        ttk.Label(
            frame,
            text="История операций",
            font=("Helvetica", 16, "bold"),
            style="primary.TLabel" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=0, column=0, pady=(0, 20), sticky="n")

        btn_frame = ttk.Frame(frame)
        btn_frame.grid(row=1, column=0, sticky="ew", pady=10)
        btn_frame.grid_columnconfigure((0, 1, 2), weight=1)

        buttons = [
            ("Обновить историю", self.show_history, "primary-outline"),
            ("Экспорт в Excel", self.export_history, "success-outline"),
            ("Удалить выбранные", self.delete_selected_history, "danger-outline"),
            ("Удалить за период", self.delete_history_by_period, "danger-outline")
        ]

        for i, (text, command, style) in enumerate(buttons):
            style = f"{style}.TButton" if 'ttkbootstrap' in sys.modules else None
            btn = ttk.Button(
                btn_frame,
                text=text,
            command=command,
                style=style,
                padding=8
            )
            btn.grid(row=0, column=i, padx=5, sticky="ew")
            ToolTip(btn, text=text, bootstyle="inverse-primary")

        columns = [
            ("id", "ID", 30),
            ("last_name", "Фамилия", 80),
            ("first_name", "Имя", 80),
            ("phone", "Телефон", 80),
            ("item_type", "Тип залога", 100),
            ("item_info", "Описание", 150),
            ("cell_number", "Ячейка", 30),
            ("created_at", "Дата приёма", 150),
            ("released_at", "Дата возврата", 150),
            ("duration", "Время нахождения", 100)
        ]

        tree_frame = ttk.Frame(frame)
        tree_frame.grid(row=2, column=0, sticky="nsew")
        tree_frame.grid_rowconfigure(0, weight=1)
        tree_frame.grid_columnconfigure(0, weight=1)

        self.tree_history = ttk.Treeview(
            tree_frame,
            columns=[col[0] for col in columns],
            show="headings",
            style="primary.Treeview" if 'ttkbootstrap' in sys.modules else None
        )

        for col in columns:
            self.tree_history.heading(col[0], text=col[1], 
                                    command=lambda c=col[0]: self.sort_treeview(self.tree_history, c))
            self.tree_history.column(col[0], width=col[2])

        def resize_columns(event):
            total_width = self.tree_history.winfo_width()
            num_columns = len(columns)
            for col, _, _ in columns:
                self.tree_history.column(col, width=max(50, int(total_width / num_columns)))

        self.tree_history.bind("<Configure>", resize_columns)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree_history.yview)
        self.tree_history.configure(yscrollcommand=scrollbar.set)
        self.tree_history.grid(row=0, column=0, sticky="nsew", pady=10)
        scrollbar.grid(row=0, column=1, sticky="ns")

        self.show_history()
        logger.debug("Вкладка 'История' настроена")

    def setup_settings_tab(self):
        logger.debug("Настройка вкладки 'Настройки'")
        frame = self.scroll_frame_settings
        frame.configure(padding="20")
        for i in range(17):
            frame.grid_rowconfigure(i, weight=0)
        frame.grid_rowconfigure(0, weight=1)
        frame.grid_columnconfigure(0, weight=1)

        ttk.Label(
            frame,
            text="Настройки приложения",
            font=("Helvetica", 16, "bold"),
            style="primary.TLabel" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=0, column=0, pady=(0, 20), sticky="n")

        # Theme toggle
        ttk.Label(
            frame,
            text="Тема интерфейса",
            font=("Helvetica", 14, "bold"),
            style="primary.TLabel" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=1, column=0, pady=(20, 5), sticky="w")

        self.theme_var = tk.BooleanVar(value=self.get_setting('theme', 'zerom') == 'zerom')
        ttk.Checkbutton(
            frame,
            text="Темная тема (выкл - светлая)",
            variable=self.theme_var,
            command=self.toggle_dark_theme,
            style="primary.Roundtoggle.TCheckbutton" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=2, column=0, pady=10, sticky="w")
        ToolTip(
            frame.winfo_children()[-1],
            text="Переключение между темной и светлой темой",
            bootstyle="inverse-primary"
        )

        # Cell range
        ttk.Label(
        frame,
            text="Диапазон ячеек:",
            font=("Helvetica", 12)
        ).grid(row=3, column=0, pady=5, sticky="w")

        cell_range_frame = ttk.Frame(frame)
        cell_range_frame.grid(row=4, column=0, sticky="ew", pady=10)
        cell_range_frame.grid_columnconfigure((1, 3), weight=1)

        ttk.Label(cell_range_frame, text="От:", font=("Helvetica", 12)).grid(row=0, column=0, padx=10, sticky="w")
        self.cell_from_var = tk.StringVar()
        self.cell_from_entry = ttk.Entry(
            cell_range_frame,
            textvariable=self.cell_from_var,
            width=10,
            font=("Helvetica", 12),
            style="primary.TEntry" if 'ttkbootstrap' in sys.modules else None
        )
        self.cell_from_entry.grid(row=0, column=1, padx=10, sticky="ew")
        self.cell_from_var.set(self.get_setting('cell_from', self.settings['cell_from']))

        ttk.Label(cell_range_frame, text="До:", font=("Helvetica", 12)).grid(row=0, column=2, padx=10, sticky="w")
        self.cell_to_var = tk.StringVar()
        self.cell_to_entry = ttk.Entry(
            cell_range_frame,
            textvariable=self.cell_to_var,
            width=10,
            font=("Helvetica", 12),
            style="primary.TEntry" if 'ttkbootstrap' in sys.modules else None
        )
        self.cell_to_entry.grid(row=0, column=3, padx=10, sticky="ew")
        self.cell_to_var.set(self.get_setting('cell_to', self.settings['cell_to']))

        ttk.Button(
            frame,
            text="Сохранить диапазон",
            command=self.save_cell_range,
            style="primary.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=8
        ).grid(row=5, column=0, pady=10, sticky="ew", padx=100)

        # Item types
        ttk.Label(
            frame,
            text="Управление типами залогов:",
            font=("Helvetica", 12)
        ).grid(row=6, column=0, pady=10, sticky="w")

        add_frame = ttk.Frame(frame)
        add_frame.grid(row=7, column=0, sticky="ew", pady=5)
        add_frame.grid_columnconfigure(1, weight=1)

        ttk.Label(add_frame, text="Новый тип:", font=("Helvetica", 12)).grid(row=0, column=0, padx=10, sticky="w")
        self.new_type_entry = ttk.Entry(
            add_frame,
            font=("Helvetica", 12),
            style="primary.TEntry" if 'ttkbootstrap' in sys.modules else None
        )
        self.new_type_entry.grid(row=0, column=1, padx=10, sticky="ew")
        ToolTip(self.new_type_entry, text="Введите новый тип залога", bootstyle="inverse-primary")

        ttk.Button(
            add_frame,
            text="Добавить",
            command=self.add_item_type,
            style="primary.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=8
        ).grid(row=0, column=2, padx=5)

        remove_frame = ttk.Frame(frame)
        remove_frame.grid(row=8, column=0, sticky="ew", pady=5)
        remove_frame.grid_columnconfigure(1, weight=1)

        ttk.Label(remove_frame, text="Удалить тип:", font=("Helvetica", 12)).grid(row=0, column=0, padx=10, sticky="w")
        self.type_to_remove = ttk.Combobox(
        remove_frame,
            font=("Helvetica", 12),
            style="primary.TCombobox" if 'ttkbootstrap' in sys.modules else None
        )
        self.type_to_remove.grid(row=0, column=1, padx=10, sticky="ew")

        ttk.Button(
            remove_frame,
            text="Удалить",
            command=self.remove_item_type,
            style="danger.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=8
        ).grid(row=0, column=2, padx=5)

        # New Sync settings
        ttk.Label(
            frame,
            text="Режим работы с базой данных",
            font=("Helvetica", 14, "bold"),
            style="primary.TLabel" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=17, column=0, pady=(20, 5), sticky="w")

        mode_frame = ttk.Frame(frame)
        mode_frame.grid(row=18, column=0, sticky="ew", pady=5)

        self.db_mode_var = tk.StringVar(value=self.get_setting('database_mode', 'server'))
        ttk.Radiobutton(
            mode_frame,
            text="Серверная база (синхронизация)",
            variable=self.db_mode_var,
            value='server',
            command=self.update_db_settings
        ).pack(side="left", padx=10)

        ttk.Radiobutton(
            mode_frame,
            text="Локальная/сетевая база",
            variable=self.db_mode_var,
            value='local',
            command=self.update_db_settings
        ).pack(side="left", padx=10)

        self.local_db_frame = ttk.Frame(frame)
        self.local_db_frame.grid(row=19, column=0, sticky="ew", pady=5)
        self.local_db_frame.grid_columnconfigure(1, weight=1)

        ttk.Label(self.local_db_frame, text="Путь к базе:").grid(row=0, column=0, padx=10, sticky="w")
        self.local_db_entry = ttk.Entry(self.local_db_frame, font=("Helvetica", 12))
        self.local_db_entry.grid(row=0, column=1, padx=10, sticky="ew")
        self.local_db_entry.insert(0, self.get_setting('local_db_path', self.settings['local_db_path']))

        ttk.Button(
            self.local_db_frame,
            text="Обзор...",
            command=self.browse_local_db,
            style="primary.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=8
        ).grid(row=0, column=2, padx=5)

        # Обновляем видимость настроек
        self.update_db_settings()
        
        # Sync settings
        ttk.Label(
            frame,
            text="Настройки синхронизации",
            font=("Helvetica", 14, "bold"),
            style="primary.TLabel" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=9, column=0, pady=(20, 5), sticky="w")

        sync_frame = ttk.Frame(frame)
        sync_frame.grid(row=10, column=0, sticky="ew", pady=5)
        sync_frame.grid_columnconfigure(1, weight=1)

        ttk.Label(sync_frame, text="Сервер:", font=("Helvetica", 12)).grid(row=0, column=0, padx=10, sticky="w")
        self.server_entry = ttk.Entry(
            sync_frame,
            font=("Helvetica", 12),
            style="primary.TEntry" if 'ttkbootstrap' in sys.modules else None
        )
        self.server_entry.grid(row=0, column=1, padx=10, sticky="ew")
        self.server_entry.insert(0, self.get_setting('server_url', self.settings['server_url']))
        ToolTip(self.server_entry, text="Введите URL сервера для синхронизации", bootstyle="inverse-primary")

        ttk.Button(
            sync_frame,
            text="Тест подключения",
            command=self.test_connection,
            style="primary.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=8
        ).grid(row=0, column=2, padx=5)

        ttk.Button(
            frame,
            text="Синхронизировать сейчас",
            command=self.sync_data,
            style="primary.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=8
        ).grid(row=11, column=0, pady=10, sticky="ew", padx=100)

        self.auto_sync_var = tk.BooleanVar(value=self.get_setting('auto_sync', 'True') == 'True')
        ttk.Checkbutton(
            frame,
            text="Автоматическая синхронизация (каждые 5 мин)",
            variable=self.auto_sync_var,
            command=self.toggle_auto_sync,
            style="primary.Roundtoggle.TCheckbutton" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=12, column=0, pady=10, sticky="w")

        # Sync interval
        sync_interval_frame = ttk.Frame(frame)
        sync_interval_frame.grid(row=13, column=0, sticky="ew", pady=5)
        sync_interval_frame.grid_columnconfigure(1, weight=1)

        ttk.Label(sync_interval_frame, text="Интервал синхронизации (минуты):", font=("Helvetica", 12)).grid(row=0, column=0, padx=10, sticky="w")
        self.sync_interval_var = tk.StringVar(value=self.get_setting('sync_interval', self.settings['sync_interval']))
        self.sync_interval_combobox = ttk.Combobox(
            sync_interval_frame,
            textvariable=self.sync_interval_var,
            values=[1, 5, 10, 30],
            width=10,
            font=("Helvetica", 12),
            style="primary.TCombobox" if 'ttkbootstrap' in sys.modules else None
        )
        self.sync_interval_combobox.grid(row=0, column=1, padx=10, sticky="ew")
        ToolTip(self.sync_interval_combobox, text="Выберите интервал автосинхронизации", bootstyle="inverse-primary")

        ttk.Button(
            sync_interval_frame,
            text="Сохранить",
            command=self.save_sync_interval,
            style="primary.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=8
        ).grid(row=0, column=2, padx=5)

        # Sync on change
        self.sync_on_change_var = tk.BooleanVar(value=self.get_setting('sync_on_change', 'False') == 'True')
        ttk.Checkbutton(
            frame,
            text="Синхронизировать при каждом изменении базы данных",
            variable=self.sync_on_change_var,
            command=self.toggle_sync_on_change,
            style="primary.Roundtoggle.TCheckbutton" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=14, column=0, pady=10, sticky="w")
        ToolTip(
            frame.winfo_children()[-1],
            text="Автоматически синхронизировать при любых изменениях в базе данных",
            bootstyle="inverse-primary"
        )

        # Template settings
        ttk.Label(
            frame,
            text="Настройки шаблона договора",
            font=("Helvetica", 14, "bold"),
            style="primary.TLabel" if 'ttkbootstrap' in sys.modules else None
        ).grid(row=15, column=0, pady=(20, 5), sticky="w")

        template_frame = ttk.Frame(frame)
        template_frame.grid(row=16, column=0, sticky="ew", pady=5)
        template_frame.grid_columnconfigure(1, weight=1)

        ttk.Label(template_frame, text="Шаблон договора:", font=("Helvetica", 12)).grid(row=0, column=0, padx=10, sticky="w")
        self.template_entry = ttk.Entry(
            template_frame,
            font=("Helvetica", 12),
            style="primary.TEntry" if 'ttkbootstrap' in sys.modules else None
        )
        self.template_entry.grid(row=0, column=1, padx=10, sticky="ew")
        self.template_entry.insert(0, self.get_setting('contract_template', self.settings['contract_template']))

        ttk.Button(
            template_frame,
            text="Обзор...",
            command=self.browse_template,
            style="primary.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=8
        ).grid(row=0, column=2, padx=5)

        ttk.Button(
            template_frame,
            text="Сохранить",
            command=lambda: self.set_setting('contract_template', self.template_entry.get()),
            style="primary.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=8
        ).grid(row=0, column=3, padx=5)
        logger.debug("Вкладка 'Настройки' настроена")

    def save_sync_interval(self):
        logger.debug("Сохранение интервала синхронизации")
        try:
            interval = int(self.sync_interval_var.get())
            if interval not in [1, 5, 10, 30]:
                raise ValueError("Недопустимый интервал синхронизации")
            self.set_setting('sync_interval', str(interval))
            logger.info(f"Интервал синхронизации сохранен: {interval} минут")
            messagebox.showinfo("Успех", f"Интервал синхронизации установлен: {interval} минут")
        except ValueError as e:
            logger.error(f"Ошибка установки интервала: {e}")
            messagebox.showerror("Ошибка", f"Недопустимый интервал: {e}")
        except Exception as e:
            logger.error(f"Ошибка сохранения интервала: {e}")
            messagebox.showerror("Ошибка", f"Ошибка при сохранении: {e}")

    def toggle_sync_on_change(self):
        logger.debug("Переключение синхронизации при изменении")
        try:
            self.sync_on_change_enabled = self.sync_on_change_var.get()
            self.set_setting('sync_on_change', str(self.sync_on_change_enabled))
            logger.info(f"Синхронизация при изменении {'включена' if self.sync_on_change_enabled else 'выключена'}")
            messagebox.showinfo("Настройки", 
                f"Синхронизация при изменении {'включена' if self.sync_on_change_enabled else 'выключена'}")
        except Exception as e:
            logger.error(f"Ошибка переключения синхронизации при изменении: {e}")
            messagebox.showerror("Ошибка", f"Не удалось изменить настройки: {e}")

    def sync_on_change(self):
        logger.debug("Проверка синхронизации при изменении")
        if self.get_setting('sync_on_change', 'False') == 'True':
            logger.debug("Запуск синхронизации при изменении")
            self.sync_data()

    def browse_template(self):
        logger.debug("Открытие диалога выбора шаблона договора")
        try:
            filename = filedialog.askopenfilename(
                title="Выберите шаблон договора",
                filetypes=[("Word Documents", "*.docx"), ("All files", "*.*")]
            )
            if filename:
                self.template_entry.delete(0, tk.END)
                self.template_entry.insert(0, filename)
                logger.info(f"Выбран шаблон договора: {filename}")
        except Exception as e:
            logger.error(f"Ошибка при выборе шаблона: {e}")
            messagebox.showerror("Ошибка", f"Не удалось выбрать файл: {e}")

    def print_contract(self):
        logger.debug("Запуск печати договора")
        selected = self.tree_search.selection()
        if not selected:
            logger.warning("Не выбран клиент для печати договора")
            messagebox.showwarning("Ошибка", "Выберите клиента для печати договора!")
            return
            
        client_data = self.tree_search.item(selected[0])['values']
        
        try:
            from docxtpl import DocxTemplate
            from docx import Document
            logger.debug("Импортированы библиотеки для работы с Word")
            
            template_path = self.get_setting('contract_template')
            logger.debug(f"Путь к шаблону договора: {template_path}")
            
            if not template_path:
                logger.error("Путь к шаблону договора не указан")
                messagebox.showerror("Ошибка", "Путь к шаблону договора не указан в настройках!")
                return
                
            if not os.path.exists(template_path):
                logger.error(f"Файл шаблона не найден: {template_path}")
                messagebox.showerror("Ошибка", f"Файл шаблона не найден по пути:\n{template_path}")
                return
            
            try:
                Document(template_path)
                logger.debug("Шаблон договора успешно открыт")
            except Exception as e:
                logger.error(f"Файл шаблона поврежден: {e}")
                messagebox.showerror("Ошибка", 
                    f"Файл шаблона поврежден или имеет неверный формат:\n{str(e)}\n"
                    "Попробуйте открыть его в Word и сохранить заново.")
                return
            
            now = datetime.now()
            logger.debug(f"Текущая дата для договора: {now}")
            
            context = {
                'номер_договора': client_data[0],
                'дата': now.strftime("%d.%m.%Y"),
                'дата_прописью': self.date_to_words(now),
                'день': now.strftime("%d"),
                'месяц': now.strftime("%m"),
                'месяц_прописью': self.month_to_words(now.month),
                'год': now.strftime("%Y"),
                'год_прописью': self.year_to_words(now.year),
                'фамилия': client_data[1].capitalize(),
                'имя': client_data[2].capitalize(),
                'телефон': client_data[3],
                'тип_залога': client_data[4].capitalize(),
                'тип_залога_стр': client_data[4].lower(),
                'описание': " в размере "+str(client_data[5])+" руб" if (str(client_data[4]) == 'Денежные средства' and str(client_data[5]) != '') else " ("+str(client_data[5])+")" if str(client_data[5]) != '' else '',
                'ячейка': client_data[6]
            }
            logger.debug(f"Контекст для шаблона: {context}")
            
            contracts_dir = os.path.join(os.getcwd(), "Договоры")
            os.makedirs(contracts_dir, exist_ok=True)
            logger.debug(f"Директория для договоров: {contracts_dir}")
            
            try:
                doc = DocxTemplate(template_path)
                doc.render(context)
                filename = os.path.join(
                    contracts_dir, 
                    f"Договор_залога_{client_data[0]}_{client_data[1]}_{client_data[2]}.docx"
                )
                doc.save(filename)
                logger.info(f"Договор сохранен: {filename}")
                self.open_file_for_printing(filename)
                
            except Exception as e:
                logger.error(f"Ошибка при заполнении шаблона: {e}")
                messagebox.showerror("Ошибка", 
                    f"Ошибка при заполнении шаблона:\n{str(e)}\n"
                    "Проверьте правильность тегов в шаблоне.")
                
        except ImportError as e:
            logger.error(f"Не установлены библиотеки для работы с Word: {e}")
            messagebox.showerror("Ошибка", 
                "Для работы с шаблонами Word установите библиотеки:\n"
                "pip install python-docx docxtpl")
        except Exception as e:
            logger.error(f"Неизвестная ошибка при создании договора: {e}")
            messagebox.showerror("Ошибка", f"Неизвестная ошибка при создании договора: {e}")

    def date_to_words(self, date):
        logger.debug("Преобразование даты в текстовый формат")
        try:
            day = date.day
            month = date.month
            year = date.year
            
            day_words = {
                1: 'первое', 2: 'второе', 3: 'третье', 4: 'четвертое', 5: 'пятое',
                6: 'шестое', 7: 'седьмое', 8: 'восьмое', 9: 'девятое', 10: 'десятое',
                11: 'одиннадцатое', 12: 'двенадцатое', 13: 'тринадцатое', 14: 'четырнадцатое',
                15: 'пятнадцатое', 16: 'шестнадцатое', 17: 'семнадцатое', 18: 'восемнадцатое',
                19: 'девятнадцатое', 20: 'двадцатое', 21: 'двадцать первое', 22: 'двадцать второе',
                23: 'двадцать третье', 24: 'двадцать четвертое', 25: 'двадцать пятое',
                26: 'двадцать шестое', 27: 'двадцать седьмое', 28: 'двадцать восьмое',
                29: 'двадцать девятое', 30: 'тридцатое', 31: 'тридцать первое'
            }
            
            month_words = {
                1: 'января', 2: 'февраля', 3: 'марта', 4: 'апреля', 5: 'мая', 6: 'июня',
                7: 'июля', 8: 'августа', 9: 'сентября', 10: 'октября', 11: 'ноября', 12: 'декабря'
            }
            
            def year_to_words(y):
                units = ['', 'один', 'два', 'три', 'четыре', 'пять', 'шесть', 'семь', 'восемь', 'девять']
                tens = ['', 'десять', 'двадцать', 'тридцать', 'сорок', 'пятьдесят', 
                       'шестьдесят', 'семьдесят', 'восемьдесят', 'девяносто']
                teens = ['десять', 'одиннадцать', 'двенадцать', 'тринадцать', 'четырнадцать',
                        'пятнадцать', 'шестнадцать', 'семнадцать', 'восемнадцать', 'девятнадцать']
                hundreds = ['', 'сто', 'двести', 'триста', 'четыреста', 'пятьсот',
                           'шестьсот', 'семьсот', 'восемьсот', 'девятьсот']
                
                if y < 10:
                    return units[y]
                elif 10 <= y < 20:
                    return teens[y - 10]
                elif 20 <= y < 100:
                    return tens[y // 10] + (' ' + units[y % 10] if y % 10 != 0 else '')
                elif 100 <= y < 1000:
                    return hundreds[y // 100] + (' ' + year_to_words(y % 100) if y % 100 != 0 else '')
                elif y == 1000:
                    return 'одна тысяча'
                elif 1000 < y < 2000:
                    return 'одна тысяча ' + year_to_words(y % 1000)
                elif 2000 <= y < 10000:
                    return units[y // 1000] + ' тысячи ' + year_to_words(y % 1000)
                else:
                    return str(y)
            
            year_word = year_to_words(year)
            if year % 10 in [1] and year % 100 != 11:
                year_word += ' год'
            elif year % 10 in [2, 3, 4] and year % 100 not in [12, 13, 14]:
                year_word += ' года'
            else:
                year_word += ' лет'
            
            result = f"{day_words.get(day, str(day))} {month_words.get(month, str(month))} {year_word}"
            logger.debug(f"Дата преобразована: {result}")
            return result
        except Exception as e:
            logger.error(f"Ошибка преобразования даты: {e}")
            return str(date)

    def month_to_words(self, month):
        logger.debug(f"Преобразование месяца {month} в текстовый формат")
        months = {
            1: 'января', 2: 'февраля', 3: 'марта', 4: 'апреля',
            5: 'мая', 6: 'июня', 7: 'июля', 8: 'августа',
            9: 'сентября', 10: 'октября', 11: 'ноября', 12: 'декабря'
        }
        result = months.get(month, str(month))
        logger.debug(f"Месяц преобразован: {result}")
        return result

    def year_to_words(self, year):
        logger.debug(f"Преобразование года {year} в текстовый формат")
        try:
            units = ['', 'один', 'два', 'три', 'четыре', 'пять', 'шесть', 'семь', 'восемь', 'девять']
            tens = ['', 'десять', 'двадцать', 'тридцать', 'сорок', 'пятьдесят', 
                   'шестьдесят', 'семьдесят', 'восемьдесят', 'девяносто']
            teens = ['десять', 'одиннадцать', 'двенадцать', 'тринадцать', 'четырнадцать',
                    'пятнадцать', 'шестнадцать', 'семнадцать', 'восемнадцать', 'девятнадцать']
            hundreds = ['', 'сто', 'двести', 'триста', 'четыреста', 'пятьсот',
                       'шестьсот', 'семьсот', 'восемьсот', 'девятьсот']
            
            def convert_part(num):
                if num < 10:
                    return units[num]
                elif 10 <= num < 20:
                    return teens[num - 10]
                elif 20 <= num < 100:
                    return tens[num // 10] + (' ' + units[num % 10] if num % 10 != 0 else '')
                elif 100 <= num < 1000:
                    return hundreds[num // 100] + (' ' + convert_part(num % 100) if num % 100 != 0 else '')
                return str(num)
            
            if year == 2000:
                return "двухтысячный"
            elif 2001 <= year <= 2099:
                thousands = "две тысячи"
                remainder = year - 2000
                return f"{thousands} {convert_part(remainder)}"
            elif 2100 <= year <= 2199:
                thousands = "две тысячи сто"
                remainder = year - 2100
                return f"{thousands} {convert_part(remainder)}"
            else:
                result = str(year)
                logger.debug(f"Год преобразован: {result}")
                return result
        except Exception as e:
            logger.error(f"Ошибка преобразования года: {e}")
            return str(year)

    def open_file_for_printing(self, filename):
        logger.debug(f"Открытие файла для печати: {filename}")
        try:
            pythoncom.CoInitialize()
            if os.name == 'nt':
                os.startfile(filename, 'print')
            elif sys.platform == 'darwin':
                subprocess.run(['open', filename])
            else:
                subprocess.run(['xdg-open', filename])
            logger.info(f"Файл открыт для печати: {filename}")
            messagebox.showinfo("Успех", f"Договор сохранен:\n{filename}\n\nОткрыт для печати!")
        except Exception as e:
            logger.error(f"Ошибка при открытии файла: {e}")
            messagebox.showinfo("Успех", 
                f"Договор сохранен:\n{filename}\n"
                f"Не удалось открыть автоматически: {str(e)}\n"
                "Откройте файл вручную.")
        finally:
            pythoncom.CoUninitialize()

    def toggle_auto_sync(self):
        logger.debug("Переключение автосинхронизации")
        try:
            self.auto_sync_enabled = self.auto_sync_var.get()
            self.set_setting('auto_sync', str(self.auto_sync_enabled))
            logger.info(f"Автосинхронизация {'включена' if self.auto_sync_enabled else 'выключена'}")
            if self.auto_sync_enabled:
                Thread(target=self.auto_sync, daemon=True).start()
            messagebox.showinfo(
                "Настройки",
                f"Автосинхронизация {'включена' if self.auto_sync_enabled else 'выключена'}"
            )
        except Exception as e:
            logger.error(f"Ошибка переключения автосинхронизации: {e}")
            messagebox.showerror("Ошибка", f"Не удалось изменить настройки: {e}")

    def toggle_dark_theme(self):
        logger.debug("Переключение темы интерфейса")
        try:
            new_theme = "zerom" if self.theme_var.get() else "reram"
            if 'ttkbootstrap' in sys.modules:
                self.style.theme_use(new_theme)
                self.root.configure(bg=self.style.colors.bg)
                self.set_setting('theme', new_theme)
                logger.info(f"Тема изменена на {new_theme}")
                messagebox.showinfo("Тема", f"Тема изменена на {'темную' if new_theme == 'zerom' else 'светлую'}")
            else:
                logger.warning("Переключение темы недоступно без ttkbootstrap")
                messagebox.showwarning("Предупреждение", "Переключение темы недоступно без установленного ttkbootstrap")
        except Exception as e:
            logger.error(f"Ошибка при переключении темы: {e}")
            messagebox.showerror("Ошибка", f"Не удалось переключить тему: {e}")

    def test_connection(self):
        logger.debug("Тестирование подключения к серверу")
        try:
            server_url = self.server_entry.get()
            self.set_setting('server_url', server_url)
            logger.debug(f"Тестируемый URL сервера: {server_url}")
            response = requests.get(f"{server_url}/ping", timeout=5)
            if response.status_code == 200:
                logger.info("Соединение с сервером успешно")
                messagebox.showinfo("Успех", "Соединение с сервером установлено!")
            else:
                logger.error(f"Сервер ответил с ошибкой: {response.text}")
                messagebox.showerror("Ошибка", f"Сервер ответил: {response.text}")
        except Exception as e:
            logger.error(f"Ошибка подключения к серверу: {e}")
            messagebox.showerror("Ошибка", f"Не удалось подключиться: {str(e)}")

    def sync_data(self):
        logger.debug("Запуск синхронизации данных")
        def sync_thread():
            try:
                server_url = self.get_setting('server_url')
                if not server_url:
                    logger.error("Не задан URL сервера")
                    messagebox.showerror("Ошибка", "Не задан URL сервера")
                    return
                last_sync = self.get_setting('last_sync', '1970-01-01 00:00:00')
                logger.debug(f"Последняя синхронизация: {last_sync}")
                response = requests.get(
                    f"{server_url}/get_updates",
                    params={'last_sync': last_sync},
                    timeout=10
                )
                if response.status_code == 200:
                    updates = response.json()
                    self.apply_updates(updates)
                    cursor = self.conn.cursor()
                    cursor.execute("""
                        SELECT * FROM clients WHERE datetime(updated_at) > datetime(?)
                        UNION ALL
                        SELECT * FROM history WHERE datetime(updated_at) > datetime(?)
                    """, (last_sync, last_sync))
                    our_updates = cursor.fetchall()
                    if our_updates:
                        columns = [description[0] for description in cursor.description]
                        updates_to_send = [dict(zip(columns, row)) for row in our_updates]
                        response = requests.post(
                            f"{server_url}/push_updates",
                            json={'updates': updates_to_send},
                            timeout=10
                        )
                        if response.status_code != 200:
                            raise Exception(f"Ошибка отправки: {response.text}")
                    new_sync_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    self.set_setting('last_sync', new_sync_time)
                    logger.info("Данные успешно синхронизированы")
                    messagebox.showinfo("Синхронизация", "Данные успешно синхронизированы!")
                    if self.tab_control.select() == self.tab_search._w:
                        self.show_all_active()
                    else:
                        self.show_history()
                else:
                    logger.error(f"Сервер ответил с ошибкой: {response.text}")
                    raise Exception(f"Сервер ответил: {response.text}")
            except Exception as e:
                logger.error(f"Ошибка синхронизации: {e}")
                messagebox.showerror("Ошибка синхронизации", f"Не удалось синхронизировать данные: {str(e)}")
        Thread(target=sync_thread, daemon=True).start()

    def apply_updates(self, updates):
        logger.debug("Применение обновлений из сервера")
        try:
            cursor = self.conn.cursor()
            for table in ['clients', 'history']:
                if table not in updates:
                    continue
                for record in updates[table]:
                    columns = list(record.keys())
                    placeholders = ', '.join(['?'] * len(columns))
                    values = list(record.values())
                    cursor.execute(f"""
                        INSERT OR REPLACE INTO {table} ({', '.join(columns)}) 
                        VALUES ({placeholders})
                    """, values)
            self.conn.commit()
            logger.info("Обновления успешно применены")
        except Exception as e:
            logger.error(f"Ошибка при применении обновлений: {e}")
            raise

    def get_setting(self, key, default=None):
        logger.debug(f"Получение настройки: {key}")
        try:
            cursor = self.conn.cursor()
            cursor.execute('SELECT value FROM settings WHERE key = ?', (key,))
            result = cursor.fetchone()
            value = result[0] if result else default
            logger.debug(f"Настройка {key}: {value}")
            return value
        except Exception as e:
            logger.error(f"Ошибка получения настройки {key}: {e}")
            return default

    def set_setting(self, key, value):
        logger.debug(f"Установка настройки: {key} = {value}")
        try:
            cursor = self.conn.cursor()
            cursor.execute('''
            INSERT OR REPLACE INTO settings (key, value)
            VALUES (?, ?)''', (key, value))
            self.conn.commit()
            logger.info(f"Настройка {key} успешно установлена")
        except Exception as e:
            logger.error(f"Ошибка установки настройки {key}: {e}")
            raise

    def load_settings(self):
        logger.debug("Загрузка настроек")
        try:
            cursor = self.conn.cursor()
            cursor.execute('SELECT value FROM settings WHERE key = "cell_from"')
            result = cursor.fetchone()
            if result:
                self.cell_from_var.set(result[0])
            cursor.execute('SELECT value FROM settings WHERE key = "cell_to"')
            result = cursor.fetchone()
            if result:
                self.cell_to_var.set(result[0])
            cursor.execute("SELECT DISTINCT item_type FROM clients ORDER BY item_type")
            types = [row[0] for row in cursor.fetchall()]
            self.type_to_remove['values'] = types
            # Добавь эти строки
            self.sync_interval_var.set(self.get_setting('sync_interval', self.settings['sync_interval']))
            self.sync_on_change_var.set(self.get_setting('sync_on_change', 'False') == 'True')
            logger.info("Настройки успешно загружены")
        except sqlite3.Error as e:
            logger.error(f"Ошибка загрузки настроек: {e}")
            messagebox.showerror("Ошибка", f"Ошибка загрузки настроек: {e}")

    def save_cell_range(self):
        logger.debug("Сохранение диапазона ячеек")
        try:
            cell_from = int(self.cell_from_var.get())
            cell_to = int(self.cell_to_var.get())
            if cell_from < 1 or cell_to < cell_from:
                raise ValueError("Некорректный диапазон ячеек: 'От' должно быть больше 0 и меньше или равно 'До'")
            cursor = self.conn.cursor()
            cursor.execute('''
                SELECT cell_number FROM clients 
                WHERE cell_number < ? OR cell_number > ?
            ''', (cell_from, cell_to))
            occupied_cells = cursor.fetchall()
            if occupied_cells:
                cells_list = ", ".join(str(cell[0]) for cell in occupied_cells)
                raise ValueError(f"Нельзя изменить диапазон: ячейки {cells_list} заняты")
            cursor.execute('''
            INSERT OR REPLACE INTO settings (key, value)
            VALUES (?, ?)''', ('cell_from', str(cell_from)))
            cursor.execute('''
            INSERT OR REPLACE INTO settings (key, value)
            VALUES (?, ?)''', ('cell_to', str(cell_to)))
            cursor.execute('DELETE FROM cells')
            for cell in range(cell_from, cell_to + 1):
                cursor.execute('INSERT INTO cells (cell_number) VALUES (?)', (cell,))
            self.conn.commit()
            logger.info(f"Диапазон ячеек сохранен: от {cell_from} до {cell_to}")
            messagebox.showinfo("Успех", f"Диапазон ячеек успешно изменен: от {cell_from} до {cell_to}!")
        except ValueError as e:
            logger.error(f"Ошибка диапазона ячеек: {e}")
            messagebox.showerror("Ошибка", str(e))
        except sqlite3.Error as e:
            logger.error(f"Ошибка сохранения диапазона ячеек: {e}")
            messagebox.showerror("Ошибка", f"Ошибка при сохранении настроек: {e}")

    def add_item_type(self):
        logger.debug("Добавление нового типа залога")
        new_type = self.new_type_entry.get().strip()
        if not new_type:
            logger.warning("Попытка добавить пустой тип залога")
            messagebox.showwarning("Ошибка", "Введите название типа")
            return
        try:
            current_types = list(self.entries['item_type']['values'])
            if new_type not in current_types:
                current_types.append(new_type)
                current_types.sort()
                self.entries['item_type']['values'] = current_types
                self.new_type_entry.delete(0, tk.END)
                logger.info(f"Добавлен тип залога: {new_type}")
                messagebox.showinfo("Успех", "Тип залога добавлен!")
            current_remove_types = list(self.type_to_remove['values'])
            if new_type not in current_remove_types:
                current_remove_types.append(new_type)
                current_remove_types.sort()
                self.type_to_remove['values'] = current_remove_types
        except Exception as e:
            logger.error(f"Ошибка добавления типа залога: {e}")
            messagebox.showerror("Ошибка", f"Ошибка при добавлении типа: {e}")

    def remove_item_type(self):
        logger.debug("Удаление типа залога")
        type_to_remove = self.type_to_remove.get()
        if not type_to_remove:
            logger.warning("Не выбран тип для удаления")
            messagebox.showwarning("Ошибка", "Выберите тип для удаления")
            return
        try:
            cursor = self.conn.cursor()
            cursor.execute('SELECT COUNT(*) FROM clients WHERE item_type = ?', (type_to_remove,))
            count = cursor.fetchone()[0]
            if count > 0:
                logger.warning(f"Попытка удалить тип '{type_to_remove}' с {count} активными залогами")
                messagebox.showwarning("Ошибка", 
                    f"Нельзя удалить тип '{type_to_remove}', так как есть {count} активных залогов этого типа")
                return
            current_types = list(self.entries['item_type']['values'])
            if type_to_remove in current_types:
                current_types.remove(type_to_remove)
                self.entries['item_type']['values'] = current_types
            current_remove_types = list(self.type_to_remove['values'])
            if type_to_remove in current_remove_types:
                current_remove_types.remove(type_to_remove)
                self.type_to_remove['values'] = current_remove_types
                self.type_to_remove.set('')
            logger.info(f"Тип залога '{type_to_remove}' удален")
            messagebox.showinfo("Успех", f"Тип '{type_to_remove}' удален из списка")
        except sqlite3.Error as e:
            logger.error(f"Ошибка удаления типа залога: {e}")
            messagebox.showerror("Ошибка", f"Ошибка при удалении типа: {e}")

    def on_tab_changed(self, event):
        logger.debug("Переключение вкладки")
        try:
            tab = self.tab_control.tab(self.tab_control.select(), "text")
            logger.debug(f"Выбрана вкладка: {tab}")
            if tab == " Поиск/Удаление ":
                self.show_all_active()
            elif tab == " История ":
                self.show_history()
            elif tab == " Настройки ":
                self.load_settings()
        except Exception as e:
            logger.error(f"Ошибка при переключении вкладки: {e}")
            messagebox.showerror("Ошибка", f"Ошибка при переключении вкладки: {e}")

    def add_client(self):
        logger.debug("Добавление нового клиента")
        try:
            data = {
                'last_name': self.entries['last_name'].get().strip(),
                'first_name': self.entries['first_name'].get().strip(),
                'phone': self.entries['phone'].get().strip(),
                'item_type': self.entries['item_type'].get().strip(),
                'item_info': self.entries['item_info'].get().strip() or ""
            }
            logger.debug(f"Данные клиента: {data}")
            if not all([data['last_name'], data['first_name'], data['phone'], data['item_type']]):
                logger.warning("Не заполнены обязательные поля")
                messagebox.showwarning("Ошибка", "Заполните все обязательные поля (помеченные *)!")
                return
            cursor = self.conn.cursor()
            cursor.execute('SELECT cell_number FROM cells WHERE is_occupied = FALSE LIMIT 1')
            free_cell = cursor.fetchone()
            if not free_cell:
                logger.warning("Нет свободных ячеек")
                messagebox.showwarning("Ошибка", "Нет свободных ячеек!")
                return
            cell_number = free_cell[0]
            created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            cursor.execute('''
            INSERT INTO clients (
                last_name, first_name, phone, item_type, item_info, cell_number, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?)''', (
                data['last_name'], data['first_name'], data['phone'],
                data['item_type'], data['item_info'], cell_number, created_at
            ))
            cursor.execute('UPDATE cells SET is_occupied = TRUE WHERE cell_number = ?', (cell_number,))
            self.conn.commit()
            self.sync_on_change()
            logger.info(f"Клиент добавлен, ячейка: {cell_number}")
            messagebox.showinfo("Успех", f"Клиент добавлен. Ячейка: {cell_number}")
            self.update_combobox()
            for name, entry in self.entries.items():
                if name != 'item_type' and isinstance(entry, ttk.Entry):
                    entry.delete(0, tk.END)
        except sqlite3.Error as e:
            logger.error(f"Ошибка добавления клиента: {e}")
            messagebox.showerror("Ошибка", f"Ошибка при добавлении: {e}")

    def show_all_active(self):
        logger.debug("Отображение всех активных клиентов")
        try:
            cursor = self.conn.cursor()
            cursor.execute('SELECT * FROM clients ORDER BY created_at DESC')
            self.display_results(self.tree_search, cursor)
            logger.info("Список активных клиентов обновлен")
        except sqlite3.Error as e:
            logger.error(f"Ошибка загрузки активных клиентов: {e}")
            messagebox.showerror("Ошибка", f"Ошибка при загрузке данных: {e}")

    def find_client(self):
        logger.debug("Поиск клиента")
        search_term = self.search_entry.get().strip()
        logger.debug(f"Поисковой запрос: {search_term}")
        if not search_term:
            self.show_all_active()
            return
        try:
            cursor = self.conn.cursor()
            cursor.execute('''
            SELECT * FROM clients 
            WHERE last_name LIKE ? OR first_name LIKE ? 
            OR phone LIKE ? OR phone LIKE ? OR cell_number = ?
            ORDER BY created_at DESC''', 
            (f'%{search_term}%', f'%{search_term}%', 
             f'%{search_term}%', f'%{search_term[-4:]}%',
             search_term))
            self.display_results(self.tree_search, cursor)
            logger.info("Поиск клиентов выполнен")
        except sqlite3.Error as e:
            logger.error(f"Ошибка поиска клиентов: {e}")
            messagebox.showerror("Ошибка", f"Ошибка поиска: {e}")

    def delete_client(self):
        logger.debug("Удаление клиента (возврат залога)")
        selected = self.tree_search.selection()
        if not selected:
            logger.warning("Не выбран клиент для удаления")
            messagebox.showwarning("Ошибка", "Выберите клиента!")
            return
        try:
            client_data = self.tree_search.item(selected[0])['values']
            cell_number = client_data[6]
            id_deleter = client_data[0]
            created = datetime.strptime(client_data[7], "%Y-%m-%d %H:%M:%S")
            logger.debug(f"Удаление клиента, ячейка: {cell_number}, ID: {id_deleter}")
            if not messagebox.askyesno("Подтверждение", f"Вернуть залог из ячейки {cell_number}?"):
                logger.info("Удаление клиента отменено")
                return
            cursor = self.conn.cursor()
            released_at = self.get_current_time().strftime("%Y-%m-%d %H:%M:%S")
            released = datetime.strptime(released_at, "%Y-%m-%d %H:%M:%S")
            if released < created:
                duration_str = "Ошибка: дата возврата раньше даты приёма"
                logger.error(duration_str)
            else:
                duration = released - created
                total_seconds = duration.total_seconds()
                days = duration.days
                hours = int((total_seconds - days * 86400) // 3600)
                minutes = int((total_seconds - days * 86400 - hours * 3600) // 60)
                seconds = int(total_seconds - days * 86400 - hours * 3600 - minutes * 60)
                duration_str = ""
                if days > 0:
                    duration_str += f"{days} дн. "
                if hours > 0 or days > 0:
                    duration_str += f"{hours} ч. "
                if minutes > 0 or hours > 0 or days > 0:
                    duration_str += f"{minutes} мин. "
                duration_str += f"{seconds} сек."
                logger.debug(f"Продолжительность залога: {duration_str}")
            cursor.execute("INSERT INTO history SELECT *, ? FROM clients WHERE cell_number = ?", ("1970-01-01 00:00:00", cell_number))
            cursor.execute("UPDATE history SET released_at = ? WHERE cell_number = ? AND ID = ?", (released_at, cell_number, id_deleter))
            cursor.execute("UPDATE history SET updated_at = ? WHERE cell_number = ? AND ID = ?", (duration_str, cell_number, id_deleter))
            cursor.execute('DELETE FROM clients WHERE cell_number = ?', (cell_number,))
            cursor.execute('UPDATE cells SET is_occupied = FALSE WHERE cell_number = ?', (cell_number,))
            self.conn.commit()
            self.sync_on_change()
            logger.info("Залог успешно возвращен")
            messagebox.showinfo("Успех", "Залог возвращен!")
            self.show_all_active()
        except sqlite3.Error as e:
            logger.error(f"Ошибка возврата залога: {e}")
            messagebox.showerror("Ошибка", f"Ошибка при возврате: {e}")

    def show_history(self):
        logger.debug("Отображение истории операций")
        try:
            cursor = self.conn.cursor()
            cursor.execute('SELECT * FROM history ORDER BY released_at DESC')
            self.display_results(self.tree_history, cursor)
            logger.info("История операций обновлена")
        except sqlite3.Error as e:
            logger.error(f"Ошибка загрузки истории: {e}")
            messagebox.showerror("Ошибка", f"Ошибка загрузки истории: {e}")

    def delete_history_record(self):
        logger.debug("Удаление записи из истории")
        selected = self.tree_history.selection()
        if not selected:
            logger.warning("Не выбрана запись для удаления")
            messagebox.showwarning("Ошибка", "Выберите запись для удаления!")
            return
        try:
            record_id = self.tree_history.item(selected[0])['values'][0]
            logger.debug(f"Удаление записи с ID: {record_id}")
            if not messagebox.askyesno(
                "Подтверждение", 
                "Вы уверены, что хотите удалить эту запись из истории?\nЭто действие нельзя отменить."):
                logger.info("Удаление записи отменено")
                return
            cursor = self.conn.cursor()
            cursor.execute('DELETE FROM history WHERE id = ?', (record_id,))
            self.conn.commit()
            self.sync_on_change()
            logger.info(f"Запись с ID {record_id} удалена")
            messagebox.showinfo("Успех", "Запись успешно удалена из истории!")
            self.show_history()
        except sqlite3.Error as e:
            logger.error(f"Ошибка удаления записи: {e}")
            messagebox.showerror("Ошибка", f"Ошибка при удалении записи: {e}")

    def export_history(self):
        logger.debug("Экспорт истории в Excel")
        try:
            cursor = self.conn.cursor()
            cursor.execute('SELECT * FROM history')
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "История залогов"
            ws.append([
                "ID", "Фамилия", "Имя", "Телефон", 
                "Тип залога", "Описание", "Ячейка", 
                "Дата приёма", "Дата возврата", "Продолжительность"
            ])
            for record in cursor.fetchall():
                ws.append(record[:10])
            os.makedirs("exports", exist_ok=True)
            file_path = "exports/history_export.xlsx"
            wb.save(file_path)
            logger.info(f"История экспортирована в {file_path}")
            messagebox.showinfo("Успех", f"Данные экспортированы в {file_path}")
        except ImportError as e:
            logger.error(f"Не установлена библиотека openpyxl: {e}")
            messagebox.showerror("Ошибка", "Для экспорта в Excel установите openpyxl:\npip install openpyxl")
        except Exception as e:
            logger.error(f"Ошибка экспорта: {e}")
            messagebox.showerror("Ошибка", f"Ошибка экспорта: {e}")

    def update_combobox(self):
        logger.debug("Обновление выпадающего списка типов залогов")
        try:
            cursor = self.conn.cursor()
            cursor.execute("SELECT DISTINCT item_type FROM clients ORDER BY item_type")
            types = [row[0] for row in cursor.fetchall()]
            default_types = self.settings['default_types']
            all_types = list(set(types + default_types))
            all_types.sort()
            self.entries['item_type']['values'] = all_types
            if all_types:
                self.entries['item_type'].current(0)
            logger.info("Список типов залогов обновлен")
        except sqlite3.Error as e:
            logger.error(f"Ошибка обновления типов залогов: {e}")
            messagebox.showerror("Ошибка", f"Ошибка загрузки типов: {e}")

    def display_results(self, tree, cursor):
        logger.debug("Отображение результатов в таблице")
        try:
            tree.delete(*tree.get_children())
            for row in cursor.fetchall():
                tree.insert("", "end", values=row)
            logger.debug("Результаты успешно отображены")
        except Exception as e:
            logger.error(f"Ошибка отображения результатов: {e}")
            raise

    def edit_client(self):
        logger.debug("Редактирование клиента")
        selected = self.tree_search.selection()
        if not selected:
            logger.warning("Не выбран клиент для редактирования")
            messagebox.showwarning("Ошибка", "Выберите клиента для редактирования!")
            return
        client_data = self.tree_search.item(selected[0])['values']
        logger.debug(f"Редактирование клиента с ID: {client_data[0]}")
        edit_window = ttk.Toplevel(self.root)
        edit_window.title("Редактирование клиента")
        edit_window.geometry("600x600")
        frame = ttk.Frame(edit_window, padding="20")
        frame.pack(fill="both", expand=True)
        ttk.Label(
            frame,
            text="Редактирование клиента",
            font=("Helvetica", 16, "bold"),
            style="primary.TLabel" if 'ttkbootstrap' in sys.modules else None
        ).pack(pady=(0, 20))
        fields = [
            ("Фамилия:", "last_name", client_data[1]),
            ("Имя:", "first_name", client_data[2]),
            ("Телефон:", "phone", client_data[3]),
            ("Тип залога:", "item_type", client_data[4]),
            ("Описание:", "item_info", client_data[5]),
            ("Ячейка:", "cell_number", client_data[6])
        ]
        entries = {}
        input_frame = ttk.Frame(frame)
        input_frame.pack(fill="x", pady=10)
        for i, (label, name, value) in enumerate(fields):
            ttk.Label(
                input_frame,
                text=label,
                font=("Helvetica", 12)
            ).grid(row=i, column=0, padx=10, pady=5, sticky="e")
            if name == "item_type":
                entry = ttk.Combobox(
                    input_frame,
                    font=("Helvetica", 12),
                    style="primary.TCombobox" if 'ttkbootstrap' in sys.modules else None
                )
                entry['values'] = self.entries['item_type']['values']
            else:
                entry = ttk.Entry(
                    input_frame,
                    font=("Helvetica", 12),
                    style="primary.TEntry" if 'ttkbootstrap' in sys.modules else None
                )
            entry.insert(0, value)
            entry.grid(row=i, column=1, padx=10, pady=5, sticky="ew")
            entries[name] = entry
            ToolTip(entry, text=f"Введите {label.strip(':')}", bootstyle="inverse-primary")
        input_frame.columnconfigure(1, weight=1)
        def save_changes():
            logger.debug("Сохранение изменений клиента")
            new_data = {
                'last_name': entries['last_name'].get(),
                'first_name': entries['first_name'].get(),
                'phone': entries['phone'].get(),
                'item_type': entries['item_type'].get(),
                'item_info': entries['item_info'].get(),
                'cell_number': entries['cell_number'].get(),
                'id': client_data[0]
            }
            logger.debug(f"Новые данные клиента: {new_data}")
            try:
                cursor = self.conn.cursor()
                cursor.execute('''
                    UPDATE clients 
                    SET last_name=?, first_name=?, phone=?, item_type=?, item_info=?, cell_number=?
                    WHERE id=?
                ''', (
                    new_data['last_name'], new_data['first_name'], new_data['phone'],
                    new_data['item_type'], new_data['item_info'], new_data['cell_number'],
                    new_data['id']
                ))
                self.conn.commit()
                self.sync_on_change()
                edit_window.destroy()
                self.show_all_active()
                logger.info("Данные клиента обновлены")
                messagebox.showinfo("Успех", "Данные клиента обновлены!")
            except sqlite3.Error as e:
                logger.error(f"Ошибка обновления клиента: {e}")
                messagebox.showerror("Ошибка", f"Ошибка при обновлении: {e}")
        ttk.Button(
            frame,
            text="Сохранить",
            command=save_changes,
            style="primary.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=10
        ).pack(pady=20)
        cancel_btn = ttk.Button(
            frame,
            text="Отмена",
            command=edit_window.destroy,
            style="secondary.Outline.TButton" if 'ttkbootstrap' in sys.modules else None,
            padding=10
        )
        cancel_btn.pack(pady=5)
        ToolTip(cancel_btn, text="Закрыть без сохранения", bootstyle="inverse-primary")

    def duplicate_client(self):
        logger.debug("Дублирование клиента")
        selected = self.tree_search.selection()
        if not selected:
            logger.warning("Не выбран клиент для дублирования")
            messagebox.showwarning("Ошибка", "Выберите клиента для дублирования!")
            return
        client_data = self.tree_search.item(selected[0])['values']
        logger.debug(f"Дублирование клиента с ID: {client_data[0]}")
        try:
            self.entries['last_name'].delete(0, tk.END)
            self.entries['last_name'].insert(0, client_data[1])
            self.entries['first_name'].delete(0, tk.END)
            self.entries['first_name'].insert(0, client_data[2])
            self.entries['phone'].delete(0, tk.END)
            self.entries['phone'].insert(0, client_data[3])
            self.entries['item_type'].set(client_data[4])
            if len(client_data) > 5:
                self.entries['item_info'].delete(0, tk.END)
                self.entries['item_info'].insert(0, client_data[5])
            self.tab_control.select(self.tab_add)
            logger.info("Данные клиента скопированы для дублирования")
        except Exception as e:
            logger.error(f"Ошибка дублирования клиента: {e}")
            messagebox.showerror("Ошибка", f"Ошибка при дублировании: {e}")

    def sort_treeview(self, tree, col, reverse=False):
        logger.debug(f"Сортировка таблицы по столбцу: {col}")
        try:
            data = [(tree.set(item, col), item) for item in tree.get_children('')]
            try:
                data.sort(key=lambda x: float(x[0]), reverse=reverse)
            except ValueError:
                data.sort(reverse=reverse)
            for index, (val, item) in enumerate(data):
                tree.move(item, '', index)
            tree.heading(col, command=lambda: self.sort_treeview(tree, col, not reverse))
            logger.debug("Таблица отсортирована")
        except Exception as e:
            logger.error(f"Ошибка сортировки таблицы: {e}")

    def auto_sync(self):
        logger.debug("Запуск автоматической синхронизации")
        while self.auto_sync_enabled:
            interval = int(self.get_setting('sync_interval', self.settings['sync_interval']))  # Получаем интервал в минутах
            logger.debug(f"Ожидание перед следующей синхронизацией: {interval} минут")
            time.sleep(interval * 60)  # Преобразуем минуты в секунды
            if not self.auto_sync_enabled:
                logger.debug("Автосинхронизация отключена, выход из цикла")
                break
            logger.debug("Выполнение автосинхронизации")
            self.sync_data()
        logger.debug("Автосинхронизация завершена")

    def setup_hotkeys(self):
        logger.debug("Настройка горячих клавиш")
        try:
            self.root.bind('<Control-n>', lambda e: self.tab_control.select(self.tab_add))
            self.root.bind('<Control-f>', lambda e: self.search_entry.focus())
            self.root.bind('<F5>', lambda e: self.show_all_active() if self.tab_control.select() == self.tab_search._w 
                          else self.show_history())
            logger.info("Горячие клавиши настроены")
        except Exception as e:
            logger.error(f"Ошибка настройки горячих клавиш: {e}")

if __name__ == "__main__":
    logger.info("Запуск главного цикла приложения")
    try:
        root = ttk.Tk() if 'ttkbootstrap' not in sys.modules else ttk.Window()
        app = PawnShopApp(root)
        root.mainloop()
        logger.info("Приложение завершило работу")
    except Exception as e:
        logger.error(f"Критическая ошибка при запуске приложения: {e}")
        raise
